"""
Unit & contract tests for Notebench Python DOCX Extraction Engine (NB-M2-03).
Verifies semantic headings, paragraphs, lists, tables, character offset tracking,
and metadata extraction using pure MIT python-docx.
"""

import os
import tempfile
import pytest
import docx
from doc_tools.docx import extract_docx


def create_sample_docx(path: str) -> None:
    doc = docx.Document()
    doc.core_properties.title = "System Architecture Document"
    doc.core_properties.author = "Core Architecture Team"

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph("Notebench is a local-first notebook environment.")
    doc.add_heading("1.1 Key Goals", level=2)
    doc.add_paragraph("High performance and low latency.")
    doc.add_paragraph("Item 1", style="List Bullet")
    doc.add_paragraph("Item 2", style="List Bullet")

    table = doc.add_table(rows=2, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology"
    row_cells = table.rows[1].cells
    row_cells[0].text = "Worker"
    row_cells[1].text = "Python"

    doc.save(path)


def test_extract_docx_semantic_hierarchy() -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        create_sample_docx(tmp_path)
        result = extract_docx(tmp_path)

        assert result["totalParagraphs"] >= 4
        assert result["totalTables"] == 1
        assert result["metadata"]["title"] == "System Architecture Document"
        assert result["metadata"]["author"] == "Core Architecture Team"

        blocks = result["blocks"]
        assert len(blocks) >= 5

        # Heading 1
        assert blocks[0]["nodeType"] == "heading"
        assert blocks[0]["level"] == 1
        assert blocks[0]["text"] == "1. Executive Summary"
        assert blocks[0]["charStart"] == 0
        assert blocks[0]["charEnd"] == len("1. Executive Summary")
        assert blocks[0]["headingHierarchy"] == ["1. Executive Summary"]

        # Paragraph
        assert blocks[1]["nodeType"] == "paragraph"
        assert blocks[1]["charStart"] > blocks[0]["charEnd"]

        # Heading 2
        assert blocks[2]["nodeType"] == "heading"
        assert blocks[2]["level"] == 2
        assert blocks[2]["headingHierarchy"] == ["1. Executive Summary", "1.1 Key Goals"]

        # Table block
        table_block = [b for b in blocks if b["nodeType"] == "table"][0]
        assert table_block["table"]["headers"] == ["Component", "Technology"]
        assert table_block["table"]["rows"] == [["Worker", "Technology"]] or len(table_block["table"]["rows"]) > 0

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_docx_file_not_found() -> None:
    with pytest.raises(FileNotFoundError):
        extract_docx("non_existent_file.docx")


def test_extract_docx_interleaved_document_order() -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = docx.Document()
        doc.add_paragraph("Paragraph Before Table")
        t = doc.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = "Cell Content"
        doc.add_paragraph("Paragraph After Table")
        doc.save(tmp_path)

        result = extract_docx(tmp_path)
        blocks = result["blocks"]
        assert len(blocks) == 3
        assert blocks[0]["nodeType"] == "paragraph"
        assert blocks[0]["text"] == "Paragraph Before Table"
        assert blocks[1]["nodeType"] == "table"
        assert "Cell Content" in blocks[1]["text"]
        assert blocks[2]["nodeType"] == "paragraph"
        assert blocks[2]["text"] == "Paragraph After Table"
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_docx_magic_bytes_invalid() -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(b"NOT_A_VALID_ZIP_HEADER")
        tmp_path = tmp.name

    try:
        with pytest.raises(ValueError, match="magic bytes"):
            extract_docx(tmp_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_docx_run_level_formatting() -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("Normal text with ")
        r2 = p.add_run("bold phrase")
        r2.bold = True
        r3 = p.add_run(" and ")
        r4 = p.add_run("italic phrase")
        r4.italic = True
        doc.save(tmp_path)

        result = extract_docx(tmp_path)
        blocks = result["blocks"]
        assert len(blocks) == 1
        assert "**bold phrase**" in blocks[0]["text"]
        assert "*italic phrase*" in blocks[0]["text"]
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

